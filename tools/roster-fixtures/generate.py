#!/usr/bin/env python3
"""Generates synthetic roster workbooks with independent producers for the XLSX reader tests.

All names and IDs are invented. No real roster data is ever used here.
Producers: XlsxWriter (mimics Excel byte-for-byte closely), openpyxl, LibreOffice (headless).
NOT covered: files saved by Microsoft Excel, Google Sheets or WPS Office themselves.

  python3 tools/roster-fixtures/generate.py   # writes web/tests/fixtures/roster/*
"""
import csv, io, os, shutil, subprocess, sys, tempfile, unicodedata, zipfile
import openpyxl
import xlsxwriter

OUT = os.path.join(os.path.dirname(__file__), "..", "..", "web", "tests", "fixtures", "roster")
os.makedirs(OUT, exist_ok=True)

# Invented people. Includes a name stored in NFD (decomposed) form and a leading-zero ID.
ROWS = [
    ("00123", "Nguyễn Văn An"),
    ("HE000124", "Trần Thị Bình"),
    ("0456789", unicodedata.normalize("NFD", "Lê Thị Ngọc Ánh")),
    ("HE000126", "Đặng Quốc Việt"),
]
HEADER = ["STT", "Mã sinh viên", "Họ và tên", "Lớp"]


def path(name):
    return os.path.join(OUT, name)


def xlsxwriter_basic():
    wb = xlsxwriter.Workbook(path("xlsxwriter-text-ids.xlsx"))
    ws = wb.add_worksheet("Danh sách")
    ws.merge_range("A1:D1", "DANH SÁCH LỚP SE0000 (dữ liệu giả)")  # title row above the header
    for c, h in enumerate(HEADER):
        ws.write_string(2, c, h)
    for r, (sid, name) in enumerate(ROWS):
        ws.write_number(3 + r, 0, r + 1)
        ws.write_string(3 + r, 1, sid)  # stored as text, keeps leading zeros
        ws.write_string(3 + r, 2, name)
        ws.write_string(3 + r, 3, "SE0000")
    ws.write_string(8, 1, "")  # blank row inside the data region
    ws.write_formula(9, 2, '=C4&""', None, "Nguyễn Văn An")  # formula with cached value
    wb.close()


def xlsxwriter_numeric_ids():
    # What Excel does when someone types 00123 into a General cell: the zeros are gone.
    wb = xlsxwriter.Workbook(path("xlsxwriter-numeric-ids.xlsx"))
    ws = wb.add_worksheet()
    ws.write_row(0, 0, ["MSSV", "Họ và tên"])
    ws.write_number(1, 0, 123)
    ws.write_string(1, 1, "Nguyễn Văn An")
    ws.write_number(2, 0, 9876543210123)
    ws.write_string(2, 1, "Trần Thị Bình")
    wb.close()


def xlsxwriter_inline_strings():
    # constant_memory mode writes inline strings instead of the shared string table.
    wb = xlsxwriter.Workbook(path("xlsxwriter-inline-strings.xlsx"), {"constant_memory": True})
    ws = wb.add_worksheet()
    ws.write_row(0, 0, ["MSSV", "Họ và tên"])
    for r, (sid, name) in enumerate(ROWS):
        ws.write_string(1 + r, 0, sid)
        ws.write_string(1 + r, 1, name)
    wb.close()


def xlsxwriter_rich_text():
    wb = xlsxwriter.Workbook(path("xlsxwriter-rich-text.xlsx"))
    ws = wb.add_worksheet()
    bold = wb.add_format({"bold": True})
    ws.write_row(0, 0, ["MSSV", "Họ và tên"])
    ws.write_string(1, 0, "HE000200")
    ws.write_rich_string(1, 1, "Phạm ", bold, "Thị", " Cúc")  # one cell made of several <r> runs
    wb.close()


def xlsxwriter_special_chars():
    wb = xlsxwriter.Workbook(path("xlsxwriter-special-chars.xlsx"))
    ws = wb.add_worksheet()
    ws.write_row(0, 0, ["MSSV", "Họ và tên"])
    ws.write_string(1, 0, "HE000300")
    ws.write_string(1, 1, "Võ <Tên> & \"Họ\" 'Lạ'\r\nDòng 2")  # XML entities and a CR/LF inside a cell
    ws.write_string(2, 0, "=1+1")  # looks like a formula, stored as text
    ws.write_string(2, 1, "Công thức giả")
    wb.close()


def openpyxl_hidden_first_sheet():
    wb = openpyxl.Workbook()
    hidden = wb.active
    hidden.title = "Ẩn"
    hidden.append(["MSSV", "Họ và tên"])
    hidden.append(["SAI0001", "Sheet ẩn không được đọc"])
    hidden.sheet_state = "hidden"
    visible = wb.create_sheet("Lớp")
    visible.append(["MSSV", "Họ và tên"])
    for sid, name in ROWS:
        visible.append([sid, name])
    wb.active = 1
    wb.save(path("openpyxl-hidden-first-sheet.xlsx"))


def openpyxl_text_and_number():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mã SV", "Họ tên", "Ghi chú"])
    ws.append(["00777", "Hoàng Văn Em", "text"])
    ws.append([778, "Vũ Thị Phương", "number"])
    ws["A4"] = "00779"
    ws["A4"].number_format = "@"
    ws["B4"] = "Bùi Văn Giang"
    wb.save(path("openpyxl-mixed.xlsx"))


def libreoffice_from_csv():
    """Real spreadsheet application: import a CSV, save as xlsx. IDs forced to text via the filter."""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "lo-from-csv.csv")
    with open(src, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["MSSV", "Họ và tên"])
        for sid, name in ROWS:
            w.writerow([sid, name])
    # CSV import filter options: separator 44, quote 34, UTF-8 (76), first line 1, column 1 = Text (2)
    out = subprocess.run(
        ["soffice", "--headless", "--convert-to", 'xlsx:Calc MS Excel 2007 XML', "--infilter=CSV:44,34,76,1,1/2", "--outdir", OUT, src],
        capture_output=True, text=True, timeout=180,
    )
    if out.returncode != 0:
        print(out.stderr, file=sys.stderr)
    shutil.rmtree(tmp, ignore_errors=True)


def libreoffice_general_cells():
    """Same CSV imported with the default (General) column type: LibreOffice turns 00123 into 123."""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "lo-general-cells.csv")
    with open(src, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["MSSV", "Họ và tên"])
        w.writerow(["00123", "Nguyễn Văn An"])
        w.writerow(["HE000124", "Trần Thị Bình"])
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--infilter=CSV:44,34,76,1", "--outdir", OUT, src],
        capture_output=True, text=True, timeout=180,
    )
    shutil.rmtree(tmp, ignore_errors=True)


def corrupt_files():
    good = open(path("xlsxwriter-text-ids.xlsx"), "rb").read()
    open(path("corrupt-truncated.xlsx"), "wb").write(good[: len(good) // 2])
    open(path("corrupt-garbage.xlsx"), "wb").write(bytes(range(256)) * 20)
    open(path("corrupt-not-a-workbook.xlsx"), "wb").write(zipfile.ZipFile(io.BytesIO(good)).read("[Content_Types].xml"))
    # A valid zip that is not a workbook.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    open(path("not-a-workbook.xlsx"), "wb").write(buf.getvalue())
    # Old binary .xls renamed to .xlsx (LibreOffice can write real BIFF8).
    tmp = tempfile.mkdtemp()
    csv_path = os.path.join(tmp, "legacy.csv")
    open(csv_path, "w", encoding="utf-8").write("MSSV,Họ và tên\nHE1,An\n")
    subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", tmp, csv_path], capture_output=True, timeout=180)
    shutil.copy(os.path.join(tmp, "legacy.xls"), path("legacy-binary-renamed.xlsx"))
    shutil.rmtree(tmp, ignore_errors=True)


def encrypted_file():
    import msoffcrypto
    with open(path("xlsxwriter-text-ids.xlsx"), "rb") as f:
        office = msoffcrypto.OfficeFile(f)
        office.load_key(password="test-only-password")
        with open(path("encrypted-password.xlsx"), "wb") as out:
            office.encrypt("test-only-password", out)


if __name__ == "__main__":
    for fn in [xlsxwriter_basic, xlsxwriter_numeric_ids, xlsxwriter_inline_strings, xlsxwriter_rich_text, xlsxwriter_special_chars,
               openpyxl_hidden_first_sheet, openpyxl_text_and_number, libreoffice_from_csv, libreoffice_general_cells, corrupt_files, encrypted_file]:
        fn()
        print("ok", fn.__name__)
