from __future__ import annotations

import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from pypdf import PdfReader


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    value = value.replace("\x00", " ").replace("\u00a0", " ")
    return re.sub(r"[ \t]+", " ", re.sub(r"\r\n?", "\n", value)).strip()


def natural_key(path: Path) -> tuple[int, str]:
    match = re.search(r"Session\s+(\d+)", path.stem, re.IGNORECASE)
    return (int(match.group(1)) if match else 10_000, path.name.lower())


def rels_by_id(zf: zipfile.ZipFile, rels_name: str) -> dict[str, dict[str, str]]:
    if rels_name not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(rels_name))
    relationships = {}
    for rel in root.findall("pr:Relationship", NS):
        relationships[rel.attrib["Id"]] = dict(rel.attrib)
    return relationships


def resolve_part(source_part: str, target: str) -> str:
    source = Path(source_part)
    combined = source.parent / target
    parts: list[str] = []
    for part in combined.parts:
        if part == "..":
            if parts:
                parts.pop()
        elif part not in (".", ""):
            parts.append(part)
    return "/".join(parts)


def xml_text(zf: zipfile.ZipFile, part_name: str) -> tuple[list[str], list[dict[str, str]]]:
    root = ET.fromstring(zf.read(part_name))
    texts = [clean_text(node.text) for node in root.findall(".//a:t", NS)]
    texts = [text for text in texts if text]
    visual_descriptions = []
    for node in root.findall(".//p:cNvPr", NS):
        item = {
            key: clean_text(node.attrib.get(key))
            for key in ("name", "title", "descr")
            if clean_text(node.attrib.get(key))
        }
        if item and (item.get("descr") or item.get("title")):
            visual_descriptions.append(item)
    return texts, visual_descriptions


def extract_pptx(path: Path) -> dict:
    with zipfile.ZipFile(path) as zf:
        presentation = ET.fromstring(zf.read("ppt/presentation.xml"))
        presentation_rels = rels_by_id(zf, "ppt/_rels/presentation.xml.rels")
        slide_parts = []
        for slide_id in presentation.findall(".//p:sldId", NS):
            rel_id = slide_id.attrib[f"{{{NS['r']}}}id"]
            target = presentation_rels[rel_id]["Target"]
            slide_parts.append(resolve_part("ppt/presentation.xml", target))

        slides = []
        all_text: list[str] = []
        all_urls: set[str] = set()
        note_count = 0
        visual_description_count = 0

        for number, slide_part in enumerate(slide_parts, start=1):
            texts, visual_descriptions = xml_text(zf, slide_part)
            rels_name = str(
                Path(slide_part).parent
                / "_rels"
                / f"{Path(slide_part).name}.rels"
            ).replace("\\", "/")
            slide_rels = rels_by_id(zf, rels_name)
            urls = sorted(
                {
                    rel["Target"]
                    for rel in slide_rels.values()
                    if rel.get("TargetMode") == "External"
                }
            )

            notes: list[str] = []
            for rel in slide_rels.values():
                if rel.get("Type", "").endswith("/notesSlide"):
                    notes_part = resolve_part(slide_part, rel["Target"])
                    notes, _ = xml_text(zf, notes_part)
                    notes = [
                        note
                        for note in notes
                        if note.lower() not in {"slide image", "presentation notes"}
                    ]
                    break

            note_count += int(bool(notes))
            visual_description_count += len(visual_descriptions)
            all_text.extend(texts)
            all_text.extend(notes)
            all_urls.update(urls)
            slides.append(
                {
                    "slide": number,
                    "text": texts,
                    "notes": notes,
                    "urls": urls,
                    "visual_descriptions": visual_descriptions,
                }
            )

        return {
            "path": str(path),
            "slide_count": len(slides),
            "slides_with_notes": note_count,
            "visual_description_count": visual_description_count,
            "urls": sorted(all_urls),
            "text_char_count": len("\n".join(all_text)),
            "slides": slides,
        }


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    all_text = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = clean_text(page.extract_text())
            error = None
        except Exception as exc:  # Keep the corpus usable if one page is malformed.
            text = ""
            error = f"{type(exc).__name__}: {exc}"
        pages.append({"page": number, "text": text, "error": error})
        all_text.append(text)
    return {
        "path": str(path),
        "page_count": len(pages),
        "text_char_count": len("\n".join(all_text)),
        "pages": pages,
    }


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    all_urls: set[str] = set()
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.hyperlink is None:
                    continue
                hyperlink = cell.hyperlink.target if cell.hyperlink else None
                if hyperlink:
                    all_urls.add(hyperlink)
                value = cell.value
                if not isinstance(value, (str, int, float, bool)) and value is not None:
                    value = str(value)
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": value,
                        "hyperlink": hyperlink,
                    }
                )
        sheets.append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "cells": cells,
            }
        )
    workbook.close()
    return {"path": str(path), "urls": sorted(all_urls), "sheets": sheets}


def token_set(text: str) -> set[str]:
    return set(re.findall(r"[0-9A-Za-zÀ-ỹ]+", text.lower()))


def jaccard(left: str, right: str) -> float:
    left_tokens = token_set(left)
    right_tokens = token_set(right)
    if not left_tokens and not right_tokens:
        return 1.0
    if not left_tokens or not right_tokens:
        return 0.0
    return len(left_tokens & right_tokens) / len(left_tokens | right_tokens)


def main() -> int:
    if len(sys.argv) != 4:
        print("Usage: extract_course_corpus.py MATERIAL_ROOT SYLLABUS_TEXT OUTPUT_DIR")
        return 2

    material_root = Path(sys.argv[1]).resolve()
    syllabus_path = Path(sys.argv[2]).resolve()
    output_dir = Path(sys.argv[3]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    pdf_paths = sorted(material_root.rglob("*.pdf"), key=natural_key)
    pptx_paths = sorted(material_root.rglob("*.pptx"), key=natural_key)
    xlsx_paths = sorted(material_root.rglob("*.xlsx"))

    pdfs = [extract_pdf(path) for path in pdf_paths]
    presentations = [extract_pptx(path) for path in pptx_paths]
    workbooks = [extract_xlsx(path) for path in xlsx_paths]
    syllabus = clean_text(syllabus_path.read_text(encoding="utf-8"))

    comparisons = []
    for pdf, pptx in zip(pdfs, presentations):
        pdf_text = "\n".join(page["text"] for page in pdf["pages"])
        pptx_text = "\n".join(
            text
            for slide in pptx["slides"]
            for text in slide["text"] + slide["notes"]
        )
        comparisons.append(
            {
                "session": natural_key(Path(pdf["path"]))[0],
                "pdf_pages": pdf["page_count"],
                "pptx_slides": pptx["slide_count"],
                "text_jaccard": round(jaccard(pdf_text, pptx_text), 4),
            }
        )

    corpus = {
        "material_root": str(material_root),
        "syllabus_path": str(syllabus_path),
        "syllabus": syllabus,
        "pdfs": pdfs,
        "presentations": presentations,
        "workbooks": workbooks,
        "comparisons": comparisons,
    }
    corpus_path = output_dir / "course_corpus.json"
    corpus_path.write_text(
        json.dumps(corpus, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    counts = Counter()
    counts["pdf_files"] = len(pdfs)
    counts["pdf_pages"] = sum(item["page_count"] for item in pdfs)
    counts["pptx_files"] = len(presentations)
    counts["pptx_slides"] = sum(item["slide_count"] for item in presentations)
    counts["slides_with_notes"] = sum(item["slides_with_notes"] for item in presentations)
    counts["pptx_urls"] = len(
        {url for presentation in presentations for url in presentation["urls"]}
    )
    counts["xlsx_files"] = len(workbooks)
    counts["xlsx_urls"] = len(
        {url for workbook in workbooks for url in workbook["urls"]}
    )
    counts["syllabus_chars"] = len(syllabus)

    summary_lines = [f"{key}: {value}" for key, value in counts.items()]
    summary_lines.append("")
    summary_lines.append("session,pdf_pages,pptx_slides,text_jaccard")
    summary_lines.extend(
        f"{item['session']},{item['pdf_pages']},{item['pptx_slides']},{item['text_jaccard']}"
        for item in comparisons
    )
    (output_dir / "extraction_summary.txt").write_text(
        "\n".join(summary_lines), encoding="utf-8"
    )

    print(corpus_path)
    print("\n".join(summary_lines[:10]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
